"""
Тесты высокоуровневого сценария (`pipeline`) и порта прогресса.

Сам конвейер переехал из «Дьяка» и покрыт эталоном приёмки; здесь
проверяется то, чего в «Дьяке» не было, — шов между библиотекой и
индикатором хода работы: фабрика получает общее число строк, `advance`
получает имя файла, а исключение проходит сквозь контекст-менеджер, не
подавляясь.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, Self

import pytest
from docx import Document
from openpyxl import Workbook

from chancellery.errors import UndefinedVariableError
from chancellery.pipeline import ProgressSink, generate_documents

if TYPE_CHECKING:
    from pathlib import Path
    from types import TracebackType


class RecordingSink:
    """Порт прогресса, записывающий всё, что ему сообщили."""

    def __init__(self, total: int) -> None:
        self.total = total
        self.files: list[str] = []
        self.entered = False
        self.exited_with: BaseException | None = None
        self.exit_called = False

    def __enter__(self) -> Self:
        self.entered = True
        return self

    def __exit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        traceback: TracebackType | None,
    ) -> None:
        self.exit_called = True
        self.exited_with = exc

    def advance(self, file: str) -> None:
        self.files.append(file)


def _table(path: Path, rows: list[list[str]]) -> Path:
    """Таблица с колонками ФИО и должностью."""
    wb = Workbook()
    ws = wb.active
    ws.append(['Фамилия', 'Имя', 'Должность'])
    for row in rows:
        ws.append(row)
    wb.save(str(path))
    return path


def _template(path: Path, text: str) -> Path:
    document = Document()
    document.add_paragraph(text)
    document.save(str(path))
    return path


def test_progress_factory_receives_row_count(tmp_path: Path) -> None:
    """Фабрика зовётся с общим числом строк — его знает только библиотека."""
    table = _table(
        tmp_path / 'emp.xlsx',
        [['Иванов', 'Пётр', 'директор'], ['Петрова', 'Анна', 'бухгалтер']],
    )
    template = _template(tmp_path / 'tpl.docx', 'Назначить {{ Фамилия | вн }}.')
    sinks: list[RecordingSink] = []

    def factory(total: int) -> ProgressSink:
        sink = RecordingSink(total)
        sinks.append(sink)
        return sink

    generate_documents(
        table,
        template,
        tmp_path / 'out',
        filename='{{ Фамилия }}.docx',
        progress_factory=factory,
    )

    assert len(sinks) == 1
    assert sinks[0].total == 2
    assert sinks[0].entered
    assert sinks[0].exit_called
    assert sinks[0].exited_with is None


def test_progress_receives_file_names_not_paths(tmp_path: Path) -> None:
    """`advance` получает ИМЯ файла — снаружи его не собрать (unique_filename)."""
    table = _table(
        tmp_path / 'emp.xlsx',
        [['Иванов', 'Пётр', 'директор'], ['Иванов', 'Пётр', 'бухгалтер']],
    )
    template = _template(tmp_path / 'tpl.docx', 'Назначить {{ Фамилия | вн }}.')
    sinks: list[RecordingSink] = []

    def factory(total: int) -> ProgressSink:
        sinks.append(RecordingSink(total))
        return sinks[-1]

    written = generate_documents(
        table,
        template,
        tmp_path / 'out',
        filename='{{ Фамилия }}.docx',
        progress_factory=factory,
    )

    # Однофамильцы: второму имени добавлен различитель, и порт видит именно его.
    assert sinks[0].files == [path.name for path in written]
    assert sinks[0].files[0] != sinks[0].files[1]


def test_progress_sees_exception_and_does_not_swallow_it(tmp_path: Path) -> None:
    """Авария доходит до `__exit__` порта и продолжает лететь наружу."""
    table = _table(tmp_path / 'emp.xlsx', [['Иванов', 'Пётр', 'директор']])
    template = _template(tmp_path / 'tpl.docx', 'Назначить {{ Неизвестное }}.')
    sinks: list[RecordingSink] = []

    def factory(total: int) -> ProgressSink:
        sinks.append(RecordingSink(total))
        return sinks[-1]

    with pytest.raises(UndefinedVariableError):
        generate_documents(
            table,
            template,
            tmp_path / 'out',
            progress_factory=factory,
            filename='{{ Фамилия }}.docx',
        )

    assert sinks[0].exit_called
    assert isinstance(sinks[0].exited_with, UndefinedVariableError)


def test_without_progress_factory_runs_silently(tmp_path: Path) -> None:
    """Прогресс необязателен: без фабрики работает заглушка."""
    table = _table(tmp_path / 'emp.xlsx', [['Иванов', 'Пётр', 'директор']])
    template = _template(tmp_path / 'tpl.docx', 'Назначить {{ Фамилия | вн }}.')

    written = generate_documents(
        table,
        template,
        tmp_path / 'out',
        filename='{{ Фамилия }}.docx',
    )

    assert len(written) == 1
    assert written[0].exists()


def test_filename_falls_back_to_ordinal_without_fio(
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Без колонок ФИО и без `filename` имена файлов порядковые, с предупреждением."""
    wb = Workbook()
    ws = wb.active
    ws.append(['Основание'])
    ws.append(['служебная записка'])
    table = tmp_path / 'plain.xlsx'
    wb.save(str(table))
    template = _template(tmp_path / 'tpl.docx', 'На основании {{ Основание | рд }}.')

    written = generate_documents(table, template, tmp_path / 'out')

    assert [path.name for path in written] == ['Документ_1.docx']
    assert 'порядковыми' in caplog.text
